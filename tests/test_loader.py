"""
Tests for loader.py. They use small made-up sheets, never the real workbook,
so they run in a second and check that messy input does not crash the loader.

Run:  .venv/bin/python -m pytest -q
"""
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import loader  # noqa: E402


def clean_rows(raw):
    """Apply the loader's cell cleaning to hand-written rows (like sheet_rows)."""
    rows = [tuple(loader.clean_value(v) for v in r) for r in raw]
    width = max(len(r) for r in rows)
    return [r + (np.nan,) * (width - len(r)) for r in rows]


# ---- cell cleaning ----------------------------------------------------------
def test_bloomberg_errors_and_dashes_become_nan():
    for v in ["#N/A N/A", "#N/A Invalid Field", "#N/A Field Not Applicable",
              "#N/A Invalid Security", "--", " -- ", "", None]:
        assert pd.isna(loader.clean_value(v)), v
    assert loader.clean_value(3.5) == 3.5
    assert loader.clean_value(" SLB US Equity ") == "SLB US Equity"


# ---- table sheets -------------------------------------------------------------
def test_table_ignores_side_notes_and_empty_rows():
    rows = clean_rows([
        ("ticker", "demo_set", None, "a side note"),
        ("SLB US Equity", "Y", None, None),
        (None, None, None, None),
        ("#N/A N/A", "--", None, None),
    ])
    df = loader.parse_table(rows)
    assert list(df.columns) == ["ticker", "demo_set"]
    assert len(df) == 1


def test_empty_sheet_gives_empty_frames():
    assert loader.parse_table([]).empty
    long, groups = loader.parse_wide([], 2)
    assert long.empty
    uni = loader.clean_universe(pd.DataFrame())
    assert uni.empty


def test_example_rows_dropped():
    df = pd.DataFrame({"ticker": ["SLB US Equity", "XOM US Equity"],
                       "headline": ["EXAMPLE: delete me", "real headline"]})
    assert len(loader.drop_example_rows(df)) == 1


def test_snapshot_dividend_and_exclusion_rules():
    df = pd.DataFrame({
        "ticker": ["A US Equity", "B US Equity", "C US Equity"],
        "NAME": ["A", "B", "C"],
        "PX_LAST": [10.0, np.nan, 5.0],
        "EQY_DVD_YLD_IND": [np.nan, np.nan, 2.0],
        "PE_RATIO": [np.nan, 12.0, "#N/A N/A"],
    })
    snap, raw = loader.clean_snapshot(df)
    s = snap.set_index("ticker")
    assert s.loc["A US Equity", "EQY_DVD_YLD_IND"] == 0.0 and s.loc["A US Equity", "div_assumed_zero"]
    assert pd.isna(s.loc["B US Equity", "EQY_DVD_YLD_IND"])        # unpriced: no assumption
    assert s.loc["B US Equity", "excluded"]
    assert pd.isna(s.loc["A US Equity", "PE_RATIO"])                 # PE never imputed
    assert pd.isna(raw.set_index("ticker").loc["A US Equity", "EQY_DVD_YLD_IND"])  # raw untouched


# ---- time series sheets ---------------------------------------------------------
def test_wide_sheet_blank_slots_errors_and_dashes():
    d1, d2 = pd.Timestamp("2026-01-02"), pd.Timestamp("2026-01-05")
    rows = clean_rows([
        ("AAA US Equity", None, None, None, "CCC US Equity", None),
        ("date", "PX_LAST", "date", "PX_LAST", "date", "PX_LAST"),
        (d1, 10.0, None, None, "#N/A N/A", "#N/A N/A"),
        (d2, "--", None, None, "#N/A N/A", "#N/A N/A"),
        ("#N/A N/A", 3.0, None, None, None, None),   # value without a date is dropped
    ])
    long, groups = loader.parse_wide(rows, 2, "Prices_D")
    assert list(long["ticker"].unique()) == ["AAA US Equity"]
    assert len(long) == 1 and long["value"].iloc[0] == 10.0
    status = dict(zip(groups["group"], groups["status"]))
    assert status == {0: "ok", 1: "blank slot", 2: "no data"}


# ---- supply chain tabs ----------------------------------------------------------
SUP_HDR = ("Ticker", "Company Name", "Cost Category", "%Cost Category", "%Supplier Rev",
           "Relationship Value (Q) (Mln) (USD)", "Relationship Year", "Relationship Period", "Source", "As Of Date")
CUS_HDR = ("Ticker", "Company Name", "%Company Rev", "Cost Category", "%Customer Cost Category",
           "Relationship Value (Q) (Mln) (USD)", "Relationship Year", "Relationship Period", "Source", "As Of Date")


def sc_block(main, header, data):
    """One 23-row block: main ticker, header, up to 20 data rows, blanks."""
    rows = [(main,) + (None,) * 9, header] + list(data)
    rows += [(None,) * 10] * (23 - len(rows))
    return rows


def test_sc_tabs_read_by_header_and_convert_percent(monkeypatch):
    monkeypatch.setitem(loader.LCFG, "sc_n_blocks", 3)
    cus = sc_block("AAA US Equity", CUS_HDR, [
        ("PETR4 BZ Equity", "Petrobras", 3.97, "CAPEX", 7.5, 346.0, "2026", "Q1", "Estimate (AAA US)", "6/5/2026"),
    ])
    empty = sc_block("BBB US Equity", CUS_HDR, [
        ("#N/A N/A", "#N/A Invalid Security") + ("#N/A Field Not Applicable",) * 8,
    ])
    rows = clean_rows(cus + empty)  # third block is missing entirely
    df, blocks = loader.parse_sc_tab(rows, "SC_Customers", "Customer",
                                     ["AAA US Equity", "BBB US Equity", "CCC US Equity"])
    assert len(df) == 1
    r = df.iloc[0]
    assert abs(r["exposure_pct"] - 0.0397) < 1e-12      # %Company Rev / 100
    assert abs(r["reverse_pct"] - 0.075) < 1e-12        # %Customer Cost Category / 100
    assert r["cost_type"] == "CAPEX"
    assert list(blocks["status"]) == ["ok", "empty block", "missing block"]


def test_supplier_tabs_merge_without_duplicates(monkeypatch):
    monkeypatch.setitem(loader.LCFG, "sc_n_blocks", 1)
    a = ("DNOW US Equity", "DNOW Inc", "COGS", 0.30, 1.88, 22.2, "2026", "Q1", "Estimate (DNOW US)", "7/7/2026")
    b = ("RUIS IJ Equity", "Radiant", "COGS", 0.03, 7.8, 2.2, "2026", "C2", "2026C2 CF", "7/31/2026")
    by_val, _ = loader.parse_sc_tab(clean_rows(sc_block("SLB US Equity", SUP_HDR, [a])), "SC_Suppliers", "Supplier", ["SLB US Equity"])
    by_rev, _ = loader.parse_sc_tab(clean_rows(sc_block("SLB US Equity", SUP_HDR, [b, a])), "SC_Suppliers_Rev", "Supplier", ["SLB US Equity"])
    merged, conflicts = loader.merge_suppliers(by_val, by_rev)
    assert len(merged) == 2 and conflicts == 0
    assert set(merged["found_in"]) == {"both", "by supplier rev %"}
    assert abs(merged.set_index("related_key").loc["DNOW US EQUITY", "exposure_pct"] - 0.003) < 1e-12


def test_matching_exact_fuzzy_and_never_forced():
    snap = pd.DataFrame({"ticker": ["TS US Equity", "UFPI US Equity", "COP US Equity"],
                         "NAME": ["TENARIS SA-ADR", "UFP INDUSTRIES INC", "CONOCOPHILLIPS"]})
    links = pd.DataFrame({
        "related_ticker": ["COP US Equity", "TEN IM Equity", "PPG US Equity", np.nan],
        "related_name": ["ConocoPhillips", "Tenaris SA", "PPG Industries Inc", "Some Private Co"],
    })
    links["related_key"] = links.apply(loader.related_key, axis=1)
    out, review = loader.match_related(links, snap)
    got = dict(zip(out["related_name"], out["universe_ticker"]))
    assert got["ConocoPhillips"] == "COP US Equity"
    assert got["Tenaris SA"] == "TS US Equity"           # fuzzy, accepted
    assert pd.isna(got["PPG Industries Inc"])            # similar name, NOT forced
    assert pd.isna(got["Some Private Co"])
    assert "Tenaris SA" in set(review["related_name"])   # fuzzy matches go to review


# ---- whole workbook ----------------------------------------------------------------
def test_tiny_workbook_with_missing_and_empty_sheets(tmp_path, monkeypatch):
    """A workbook with only a few sheets, some empty, must load without crashing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Universe"
    ws.append(["ticker", "chain", "role", "notes", "raw_ticker", "demo_set", "dup_check"])
    ws.append(["AAA US Equity", "Oilfield services", "Core", None, None, "Y", None])
    ws.append(["BBB US Equity", None, None, None, None, None, None])
    ws.append([None] * 7)  # blank slot
    s = wb.create_sheet("Snapshot")
    s.append(["ticker", "asof", "NAME", "GICS_SECTOR_NAME", "GICS_INDUSTRY_NAME", "PX_LAST", "EQY_DVD_YLD_IND", "BACKLOG"])
    s.append(["AAA US Equity", "2026-10-19", "AAA CORP", "Energy", "Oil", 10, "--", "#N/A Invalid Field"])
    s.append(["BBB US Equity", "2026-10-19", "BBB INC", "#N/A N/A", None, "#N/A N/A", 1.5, "#N/A Invalid Field"])
    wb.create_sheet("Prices_D")      # empty
    wb.create_sheet("SC_Customers")  # empty
    path = tmp_path / "tiny.xlsx"
    wb.save(path)

    monkeypatch.setattr(loader.config, "CACHE_DIR", tmp_path)
    d = loader.load_data(force=True, path=path)
    assert len(d["universe"]) == 2
    assert d["broken_fields"] == ["BACKLOG"]
    assert d["snapshot"].set_index("ticker").loc["BBB US Equity", "excluded"]
    assert d["prices"].empty and d["links"].empty
    loader.write_quality_report(d)  # must not crash either
    loader.data_summary(d)
