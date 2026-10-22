"""
make_sample_data.py: write a FAKE workbook in the exact Bloomberg layout.

    .venv/bin/python make_sample_data.py              # -> sample/sample_workbook.xlsx
    SCREENER_DATA_FILE=sample/sample_workbook.xlsx .venv/bin/streamlit run app.py

Use it to try the screener without a Bloomberg terminal, or as a reference
for the layout your own pull must follow (see docs/DATA_FORMAT.md).

Everything in it is invented: company names, tickers (SMPxx US Equity,
EXTxx ... Equity), prices and fundamentals come from a random number
generator with a fixed seed. It also contains the same kinds of problems as
a real values-only export (Bloomberg error strings, "--", blank slots,
broken fields, delisted tickers, empty supply chain blocks) so every
cleaning step gets exercised.
"""
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

import config

ASOF = pd.Timestamp("2026-10-19")
NA = "#N/A N/A"
rng = np.random.default_rng(42)

# (sector, industry, number of companies, name word, oil beta, copper beta, steel beta, gross margin)
INDUSTRIES = [
    ("Energy", "Oil, Gas & Consumable Fuels", 14, "ENERGY", 0.55, 0.0, 0.0, 42),
    ("Energy", "Energy Equipment & Services", 10, "DRILLING", 0.45, 0.0, 0.0, 20),
    ("Industrials", "Machinery", 14, "MACHINERY", 0.05, 0.15, 0.1, 33),
    ("Industrials", "Electrical Equipment", 10, "ELECTRIC", -0.05, 0.3, 0.0, 35),
    ("Industrials", "Aerospace & Defense", 9, "AEROSPACE", -0.1, 0.05, 0.0, 22),
    ("Industrials", "Marine Transportation", 2, "SHIPPING", 0.1, 0.0, 0.0, 30),     # too small: sector ranking
    ("Materials", "Metals & Mining", 3, "STEEL", 0.05, 0.2, 0.5, 15),              # too small: all-company ranking
]
CORE_PER_INDUSTRY = {"Oil, Gas & Consumable Fuels": 4, "Energy Equipment & Services": 3, "Machinery": 3,
                     "Electrical Equipment": 2}
CHAINS = {"Oil, Gas & Consumable Fuels": "Producers", "Energy Equipment & Services": "Oilfield services",
          "Machinery": "Equipment", "Electrical Equipment": "Grid equipment"}
PREFIXES = ["NORTHWIND", "BLUESTONE", "REDWATER", "IRONPEAK", "SILVERLINE", "GRANITE", "HARBOR", "SUMMIT",
            "PRAIRIE", "CASCADE", "COPPERFIELD", "WESTGATE", "LAKESHORE", "PINECREST", "STONEBRIDGE",
            "CLEARWATER", "HIGHLAND", "RIVERBEND", "EASTFIELD", "OAKRIDGE", "BRIGHTON", "FALCON", "MERIDIAN",
            "TIDEWAY", "ASHFORD", "BLACKROCK HILL", "CEDAR", "DUNMORE", "ELMWOOD", "FAIRHAVEN", "GLENVIEW",
            "HOLLOWAY", "JUNIPER", "KESTREL", "LANCASTER", "MAPLE", "NEWPORT", "ORCHARD", "PEMBROKE",
            "QUARRY", "ROCKPORT", "SHERWOOD", "THORNTON", "UNION BAY", "VALLEY", "WILLOW", "YARROW",
            "ZENITH", "ALDER", "BIRCH", "CORAL", "DRIFTWOOD", "EVEREST", "FOXGLOVE", "GALLOWAY", "HAWTHORN",
            "INDIGO", "JASPER", "KINGSTON", "LIBERTY", "MARLOW", "NEPTUNE", "OSPREY", "PINNACLE"]
SUFFIXES = ["CORP", "INC", "CO", "HOLDINGS INC", "LTD"]
COMMODITIES = [("CL1 Comdty", 60.0, 0.020), ("CO1 Comdty", 64.0, 0.019), ("NG1 Comdty", 3.4, 0.035),
               ("HG1 Comdty", 4.6, 0.013), ("HRC1 Comdty", 800.0, 0.012)]
SNAPSHOT_COLS = ["ticker", "asof", "NAME", "GICS_SECTOR_NAME", "GICS_INDUSTRY_NAME", "CUR_MKT_CAP", "PX_LAST",
                 "PX_VOLUME", "CHG_PCT_YTD", "CHG_PCT_1M", "SALES_GROWTH", "GROSS_MARGIN", "OPER_MARGIN",
                 "TOT_DEBT_TO_COM_EQY", "NET_DEBT_TO_EBITDA", "RETURN_COM_EQY", "RETURN_ON_ASSET", "CUR_RATIO",
                 "CF_FREE_CASH_FLOW", "FREE_CASH_FLOW_YIELD", "PE_RATIO", "PX_TO_BOOK_RATIO", "EV_TO_T12M_EBITDA",
                 "EQY_DVD_YLD_IND", "BS_INVENTORIES", "INVENTORY_TURNOVER", "BACKLOG", "CAPITAL_EXPEND",
                 "BEST_CAPEX", "EQY_REC_CONS", "BEST_TARGET_PRICE", "SHORT_INT_RATIO", "EQY_INST_PCT_SH_OUT",
                 "NEWS_SENTIMENT_DAILY_AVG", "NEWS_HEAT_PUB_AVG", "BEST_EPS_FY1", "BEST_EPS_FY2"]
BROKEN = ["INVENTORY_TURNOVER", "BACKLOG", "NEWS_HEAT_PUB_AVG"]   # 100 percent errors, like the real pull


# ---------------------------------------------------------------------------
# Companies
# ---------------------------------------------------------------------------
def make_companies():
    """One row per company: identity plus the 'true' traits used below."""
    rows, k = [], 0
    for sector, industry, n, word, b_oil, b_cu, b_steel, gm in INDUSTRIES:
        for j in range(n):
            k += 1
            rows.append(dict(
                ticker="SMP%02d US Equity" % k,
                name="%s %s %s" % (PREFIXES[(k - 1) % len(PREFIXES)], word, SUFFIXES[k % len(SUFFIXES)]),
                sector=sector, industry=industry,
                is_core=j < CORE_PER_INDUSTRY.get(industry, 0),
                chain=CHAINS.get(industry) if j < CORE_PER_INDUSTRY.get(industry, 0) else None,
                b_oil=b_oil + rng.normal(0, 0.15), b_cu=b_cu + rng.normal(0, 0.1), b_steel=b_steel + rng.normal(0, 0.1),
                gm=gm + rng.normal(0, 7), quality=rng.normal(0, 1),
                no_inventory=industry == "Oil, Gas & Consumable Fuels" and j % 2 == 1, delisted=False,
            ))
    # Two delisted tickers: in the Universe, but no price (one has a short
    # price history that stops early, one has none at all).
    for name in ["OLDCO ENERGY CORP", "MERGED MACHINERY INC"]:
        k += 1
        rows.append(dict(ticker="SMP%02d US Equity" % k, name=name, sector=None, industry=None, is_core=False,
                         chain=None, b_oil=0, b_cu=0, b_steel=0, gm=30, quality=0, no_inventory=False, delisted=True))
    comp = pd.DataFrame(rows)
    # Core names first (Universe rows 2..), like the real workbook.
    return pd.concat([comp[comp["is_core"]], comp[~comp["is_core"]]], ignore_index=True)


# ---------------------------------------------------------------------------
# Time series
# ---------------------------------------------------------------------------
def make_prices(comp, days):
    """Commodity random walks and stock prices with commodity betas."""
    n = len(days)
    cl = rng.normal(0, 0.02, n)
    comm_rets = {"CL1 Comdty": cl, "CO1 Comdty": 0.95 * cl + rng.normal(0, 0.006, n)}
    for code, _, vol in COMMODITIES[2:]:
        comm_rets[code] = rng.normal(0, vol, n)
    comm = {code: start * np.exp(np.cumsum(comm_rets[code])) for code, start, _ in COMMODITIES}
    market = rng.normal(0.0003, 0.01, n)
    prices = {}
    for _, c in comp.iterrows():
        if c["delisted"]:
            continue
        r = (market + c["b_oil"] * comm_rets["CL1 Comdty"] + c["b_cu"] * comm_rets["HG1 Comdty"]
             + c["b_steel"] * comm_rets["HRC1 Comdty"] + rng.normal(0.0002 * c["quality"], 0.015, n))
        prices[c["ticker"]] = rng.uniform(15, 250) * np.exp(np.cumsum(r))
    return comm, prices


def wide_sheet(ws, groups, width):
    """Write side-by-side column groups: row 1 ticker, row 2 headers, data from row 3.
    groups: list of (ticker or None, headers, list_of_rows)."""
    max_rows = max((len(g[2]) for g in groups), default=0)
    grid = [[None] * (width * len(groups)) for _ in range(max_rows + 2)]
    for i, (ticker, headers, rows) in enumerate(groups):
        c0 = i * width
        if ticker is None:
            continue                       # blank slot: no ticker, no data
        grid[0][c0] = ticker
        for j, h in enumerate(headers):
            grid[1][c0 + j] = h
        for r, row in enumerate(rows):
            for j, v in enumerate(row):
                grid[2 + r][c0 + j] = v
    for row in grid:
        ws.append(row)


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------
def build(path):
    comp = make_companies()
    days = pd.bdate_range("2025-10-20", ASOF)
    comm, prices = make_prices(comp, days)
    quarters = pd.date_range("2021-12-31", "2026-06-30", freq="QE")
    weeks = pd.date_range("2025-10-24", "2026-10-16", freq="W-FRI")
    news_days = pd.date_range(ASOF - pd.Timedelta(days=90), ASOF, freq="D")
    wb = openpyxl.Workbook()

    # README / Settings ------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    for line in ["SAMPLE DATA: every company, ticker and number in this workbook is invented.",
                 "Generated by make_sample_data.py in the exact layout of a Bloomberg values-only pull.",
                 "See docs/DATA_FORMAT.md for the layout."]:
        ws.append([line])
    ws = wb.create_sheet("Settings")
    ws.append(["setting", "value", "note"])
    for k, v, n in [("end_date", ASOF, "End of every history pull"),
                    ("prices_start", datetime(2025, 10, 19), "1 year of daily prices"),
                    ("fundamentals_start", datetime(2021, 10, 20), "5 years of quarterly fundamentals"),
                    ("sentiment_start", datetime(2026, 7, 21), "90 days of daily news sentiment"),
                    ("estimates_start", datetime(2025, 10, 19), "1 year of weekly consensus estimates")]:
        ws.append([k, v.to_pydatetime() if hasattr(v, "to_pydatetime") else v, n])

    # Universe -----------------------------------------------------------------
    ws = wb.create_sheet("Universe")
    ws.append(["ticker", "chain", "role", "notes", "raw_ticker", "demo_set", "dup_check", None,
               "Core names (demo_set = Y) first. Remaining rows are empty slots."])
    for _, c in comp.iterrows():
        ws.append([c["ticker"], c["chain"], "Core" if c["is_core"] else None, None, c["ticker"],
                   "Y" if c["is_core"] else None, None])
    for _ in range(6):
        ws.append([None] * 7)          # blank slots, like the real template

    # Fundamentals (quarterly) -----------------------------------------------------
    fund, snap_extra = [], {}
    for _, c in comp.iterrows():
        if c["delisted"]:
            fund.append((c["ticker"], ["date", "BS_INVENTORIES", "SALES_REV_TURN", "GROSS_MARGIN", "EV_TO_T12M_EBITDA"],
                         [["#N/A Invalid Security", None, None, None, None]]))
            continue
        sales0 = rng.lognormal(6.5, 1.0)
        growth = rng.normal(0.012, 0.02)
        sales = sales0 * np.exp(np.cumsum(rng.normal(growth, 0.05, len(quarters))))
        inv_ratio = 0 if c["no_inventory"] else rng.uniform(0.3, 1.2)
        inv = sales * inv_ratio * np.exp(np.cumsum(rng.normal(0, 0.04, len(quarters))))
        gm = c["gm"] + rng.normal(0, 1.5, len(quarters))
        ev_mean = rng.uniform(6, 18)
        ev = ev_mean * np.exp(np.cumsum(rng.normal(0, 0.07, len(quarters))))
        rows = []
        for q, s_, i_, g_, e_ in zip(quarters, sales, inv, gm, ev):
            rows.append([q.to_pydatetime(), round(i_, 1), round(s_, 1), round(g_, 3),
                         round(e_, 4) if rng.random() > 0.03 else NA])
        fund.append((c["ticker"], ["date", "BS_INVENTORIES", "SALES_REV_TURN", "GROSS_MARGIN", "EV_TO_T12M_EBITDA"], rows))
        snap_extra[c["ticker"]] = dict(inv=inv[-1], sales_ttm=sales[-4:].sum(), gm=gm[-1], ev=ev[-1] * rng.uniform(0.9, 1.1),
                                       sales_growth=100 * (sales[-4:].sum() / sales[-8:-4].sum() - 1))

    # Estimates (weekly) ----------------------------------------------------------------
    est = []
    for _, c in comp.iterrows():
        if c["delisted"]:
            est.append((c["ticker"], ["date", "BEST_EPS", "EQY_REC_CONS"], []))
            continue
        eps = abs(rng.normal(3, 2)) + 0.05
        eps_path = eps * np.exp(np.cumsum(rng.normal(0.002 * c["quality"], 0.015, len(weeks))))
        rec = np.clip(4 + rng.normal(0, 0.4) + np.cumsum(rng.choice([0, 0, 0, 0, 0.1, -0.1], len(weeks))), 1.5, 5)
        rows = [[w.to_pydatetime(), round(e, 3) if rng.random() > 0.1 else NA, round(r, 3)]
                for w, e, r in zip(weeks, eps_path, rec)]
        est.append((c["ticker"], ["date", "BEST_EPS", "EQY_REC_CONS"], rows))
        snap_extra[c["ticker"]].update(eps=eps_path[-1], rec=rec[-1])

    # Sentiment (daily, only on days with news) ------------------------------------------
    sent = []
    for _, c in comp.iterrows():
        p_news = 0.02 if c["delisted"] else rng.choice([0.04, 0.3, 0.5, 0.7])   # some get "low coverage"
        days_with_news = [d for d in news_days if d.weekday() < 5 and rng.random() < p_news]
        rows = [[d.to_pydatetime(), round(float(np.clip(rng.normal(0.05 + 0.1 * c["quality"], 0.3), -1, 1)), 4)]
                for d in days_with_news]
        sent.append((c["ticker"], ["date", "NEWS_SENTIMENT_DAILY_AVG"], rows if not c["delisted"] else []))

    # Snapshot -------------------------------------------------------------------------------
    ws = wb.create_sheet("Snapshot")
    ws.append(SNAPSHOT_COLS)
    for _, c in comp.iterrows():
        t = c["ticker"]
        row = {k: NA for k in SNAPSHOT_COLS}
        row.update(ticker=t, asof=ASOF.to_pydatetime(), NAME=c["name"])
        for b in BROKEN:
            row[b] = "#N/A Invalid Field"
        if c["delisted"]:
            ws.append([row[k] for k in SNAPSHOT_COLS])
            continue
        px = prices[t]
        x = snap_extra[t]
        q = c["quality"]
        mcap = rng.lognormal(23.0, 1.0)
        fcf_y = rng.normal(4 + q, 3)
        row.update(
            GICS_SECTOR_NAME=c["sector"], GICS_INDUSTRY_NAME=c["industry"], CUR_MKT_CAP=mcap, PX_LAST=round(px[-1], 3),
            PX_VOLUME=int(rng.lognormal(14, 1)), CHG_PCT_YTD=100 * (px[-1] / px[-200] - 1),
            CHG_PCT_1M=100 * (px[-1] / px[-22] - 1), SALES_GROWTH=x["sales_growth"], GROSS_MARGIN=x["gm"],
            OPER_MARGIN=x["gm"] * 0.45 + rng.normal(2 * q, 4), TOT_DEBT_TO_COM_EQY=rng.lognormal(4.0 - 0.2 * q, 0.6),
            NET_DEBT_TO_EBITDA=rng.normal(1.5 - 0.3 * q, 1.2), RETURN_COM_EQY=rng.normal(14 + 3 * q, 7),
            RETURN_ON_ASSET=rng.normal(6 + 1.5 * q, 3.5), CUR_RATIO=rng.lognormal(0.45, 0.35),
            CF_FREE_CASH_FLOW=mcap * fcf_y / 100 / 1e6, FREE_CASH_FLOW_YIELD=fcf_y,
            PE_RATIO=rng.lognormal(3.0, 0.4) if rng.random() > 0.1 else NA, PX_TO_BOOK_RATIO=rng.lognormal(1.1, 0.6),
            EV_TO_T12M_EBITDA=x["ev"], EQY_DVD_YLD_IND=rng.uniform(0.5, 5) if rng.random() > 0.25 else NA,
            BS_INVENTORIES=x["inv"], CAPITAL_EXPEND=-x["sales_ttm"] * rng.uniform(0.02, 0.25),
            BEST_CAPEX=-x["sales_ttm"] * rng.uniform(0.005, 0.06), EQY_REC_CONS=x["rec"],
            BEST_TARGET_PRICE=px[-1] * (1 + rng.normal(0.1, 0.12)), SHORT_INT_RATIO=rng.lognormal(1.2, 0.4),
            EQY_INST_PCT_SH_OUT=rng.uniform(55, 98), NEWS_SENTIMENT_DAILY_AVG=rng.normal(0.05, 0.2),
            BEST_EPS_FY1=x["eps"], BEST_EPS_FY2=x["eps"] * rng.uniform(1.0, 1.2),
        )
        if rng.random() < 0.05:
            row["GROSS_MARGIN"] = "--"     # Bloomberg sometimes shows "--"
        ws.append([row[k] for k in SNAPSHOT_COLS])
    # A few extreme values, so winsorizing and the caveat badge have work to do.
    ws.cell(row=3, column=SNAPSHOT_COLS.index("PE_RATIO") + 1, value=650.0)
    ws.cell(row=4, column=SNAPSHOT_COLS.index("PX_TO_BOOK_RATIO") + 1, value=420.0)
    ws.cell(row=5, column=SNAPSHOT_COLS.index("OPER_MARGIN") + 1, value=-900.0)

    # Prices, fundamentals, estimates, sentiment, commodities -----------------------------
    price_groups, n_delisted = [], 0
    for _, c in comp.iterrows():
        t = c["ticker"]
        if c["delisted"]:
            n_delisted += 1
            if n_delisted == 1:     # history stops after 60 days
                rows = [[d.to_pydatetime(), round(50 + i * 0.1, 3)] for i, d in enumerate(days[:60])]
            else:                   # nothing at all
                rows = [["#N/A Invalid Security", None]]
            price_groups.append((t, ["date", "PX_LAST"], rows))
        else:
            price_groups.append((t, ["date", "PX_LAST"], [[d.to_pydatetime(), round(p, 3)] for d, p in zip(days, prices[t])]))
    blanks = [(None, [], [])] * 5
    wide_sheet(wb.create_sheet("Prices_D"), price_groups + blanks, 2)
    wide_sheet(wb.create_sheet("Fundamentals_Q"), fund + blanks, 5)
    wide_sheet(wb.create_sheet("Estimates_W"), est + blanks, 3)
    wide_sheet(wb.create_sheet("Sentiment_D"), sent + blanks, 2)
    wide_sheet(wb.create_sheet("Commodities_D"),
               [(code, ["date", "PX_LAST"], [[d.to_pydatetime(), round(v, 3)] for d, v in zip(days, comm[code])])
                for code, _, _ in COMMODITIES], 2)

    # Links / News (template example rows, plus two invented headlines) ------------------
    ws = wb.create_sheet("Links")
    ws.append(["main_company", "direction", "related_company", "related_ticker", "pct_of_revenue_or_cost", "notes"])
    ws.append([comp.at[0, "ticker"], "Customer", "Example Energy Co", "EXAM US Equity", 0.05, "EXAMPLE ROW - delete"])
    ws = wb.create_sheet("News")
    ws.append(["ticker", "date", "headline", "manual_score"])
    ws.append([comp.at[0, "ticker"], "2026-10-15", "EXAMPLE: company raises full-year guidance - delete", 1])
    ws.append([comp.at[0, "ticker"], "2026-10-14", "Sample headline: company wins large drilling contract", 1])
    ws.append([comp.at[0, "ticker"], "2026-10-02", "Sample headline: plant outage cuts quarterly output", -1])

    # Data dictionary ------------------------------------------------------------------------
    ws = wb.create_sheet("Data_Dictionary")
    ws.append(["group", "column_name", "sheet", "update_frequency", "unit", "what_it_is", "why_it_matters",
               "how_to_read", "field_status"])
    for f in SNAPSHOT_COLS[2:]:
        ws.append(["Snapshot", f, "Snapshot", "daily", "", "", "", "", "unverified" if f in BROKEN else "verified (sample)"])

    # Supply chain tabs --------------------------------------------------------------------------
    write_supply_chain(wb, comp)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return comp


def write_supply_chain(wb, comp):
    """SC_Suppliers, SC_Suppliers_Rev, SC_Customers: one 23-row block per core name."""
    exch = ["LN", "GR", "FP", "JP", "CN", "BZ", "AB", "NO", "IM", "AU"]
    words = ["ATLANTIC", "NORDIC", "PACIFIC", "ANDEAN", "GULF", "ALPINE", "CELTIC", "IBERIAN", "BALTIC", "SAHARA"]
    kinds = ["OFFSHORE PLC", "STEEL AG", "LOGISTICS SA", "PIPE CO", "CHEMICALS NV", "ENGINEERING LTD", "ENERGY ASA",
             "RESOURCES LTD", "SERVICES SPA", "POWER CORP"]
    external = [("EXT%02d %s Equity" % (i, exch[i % 10]), ("%s %s" % (words[i % 10], kinds[(i // 10) % 10])).title())
                for i in range(1, 81)]
    # One external listing of a universe company (like a foreign line of an
    # ADR): different ticker, same name, so the fuzzy name match is exercised.
    uni_pick = comp[~comp["is_core"] & ~comp["delisted"]].iloc[0]
    external.append(("EXT99 LN Equity", uni_pick["name"].title()))
    scored = comp[~comp["delisted"]]
    core = comp[comp["is_core"]].reset_index(drop=True)
    sup_hdr = ["Ticker", "Company Name", "Cost Category", "%Cost Category", "%Supplier Rev",
               "Relationship Value (Q) (Mln) (USD)", "Relationship Year", "Relationship Period", "Source", "As Of Date"]
    cus_hdr = ["Ticker", "Company Name", "%Company Rev", "Cost Category", "%Customer Cost Category",
               "Relationship Value (Q) (Mln) (USD)", "Relationship Year", "Relationship Period", "Source", "As Of Date"]
    tabs = {name: wb.create_sheet(name) for name in ["SC_Suppliers", "SC_Suppliers_Rev", "SC_Customers"]}

    def related(main, n):
        pool_uni = [(t, nm.title()) for t, nm in zip(scored["ticker"], scored["name"]) if t != main]
        picks = []
        for _ in range(n):
            if rng.random() < 0.3:
                picks.append(pool_uni[rng.integers(len(pool_uni))])
            else:
                picks.append(external[rng.integers(len(external))])
        return list(dict(picks).items())   # drop duplicates, keep order

    for i, c in core.iterrows():
        main = c["ticker"]
        root = main.replace(" Equity", "")
        sups = []
        for t, nm in related(main, rng.integers(8, 21)):
            cat = rng.choice(["COGS"] * 15 + ["CAPEX"] * 3 + ["SGA", NA])
            src = "Estimate (%s)" % t.replace(" Equity", "") if rng.random() < 0.8 else "2025A CF (%s)" % t.replace(" Equity", "")
            sups.append([t, nm, cat, round(rng.uniform(0.01, 3), 4), round(rng.uniform(0.1, 20), 3),
                         round(rng.uniform(0.5, 300), 3), str(rng.choice([2025, 2026])), rng.choice(["Q1", "Q2", "A"]),
                         src, "%d/%d/2026" % (rng.integers(1, 10), rng.integers(1, 28))])
        by_value = sorted(sups, key=lambda r: -r[5])
        by_rev = sorted(sups, key=lambda r: -r[4])[: max(3, len(sups) - 2)]
        if i in (2, 5):   # two core names without quantified customers, like the real pull
            cus = [[NA, "#N/A Invalid Security"] + ["#N/A Field Not Applicable"] * 8]
        else:
            cus = []
            for t, nm in related(main, rng.integers(5, 16)):
                cat = rng.choice(["CAPEX"] * 6 + ["COGS"] * 6 + [NA])
                cus.append([t, nm, round(rng.uniform(0.05, 6), 4), cat, round(rng.uniform(0.1, 10), 3),
                            round(rng.uniform(1, 400), 3), "2026", "Q1", "Estimate (%s)" % root,
                            "%d/%d/2026" % (rng.integers(1, 10), rng.integers(1, 28))])
            cus.sort(key=lambda r: -r[2])
        for name, hdr, rows in [("SC_Suppliers", sup_hdr, by_value), ("SC_Suppliers_Rev", sup_hdr, by_rev),
                                ("SC_Customers", cus_hdr, cus)]:
            ws = tabs[name]
            ws.append([main, "<- main company (block %d)" % (i + 1)] + [None] * 8)
            ws.append(hdr)
            for r in rows[:20]:
                ws.append(r)
            for _ in range(20 - min(len(rows), 20)):
                ws.append([None] * 10)
            ws.append([None] * 10)          # row 23: blank separator

    ws = wb.create_sheet("SC_README")
    ws.append(["SAMPLE supply chain tabs: invented companies, layout identical to a Bloomberg SPLC pull."])
    ws = wb.create_sheet("SC_Settings")
    ws.append(["Supply chain pull settings"])
    ws.append([None])
    ws.append(["list", "Bloomberg field (BDS)", "Override string (BDS)", "Relationship letter (BDP)", "Notes"])
    ws.append(["SC_Suppliers", "SUPPLY_CHAIN_SUPPLIERS",
               "SUPPLY_CHAIN_SUM_COUNT_OVERRIDE=20,QUANTIFIED_OVERRIDE=Y,SUP_CHAIN_RELATIONSHIP_SORT_OVR=C", "S",
               "Top 20 suppliers ranked by relationship value"])
    ws.append(["SC_Suppliers_Rev", "SUPPLY_CHAIN_SUPPLIERS",
               "SUPPLY_CHAIN_SUM_COUNT_OVERRIDE=20,QUANTIFIED_OVERRIDE=Y,SUP_CHAIN_RELATIONSHIP_SORT_OVR=R", "S",
               "Top 20 suppliers ranked by the supplier's revenue share"])
    ws.append(["SC_Customers", "SUPPLY_CHAIN_CUSTOMERS", "SUPPLY_CHAIN_SUM_COUNT_OVERRIDE=20,QUANTIFIED_OVERRIDE=Y", "C",
               "Top 20 customers ranked by share of the main company's revenue"])


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else config.PROJECT_DIR / "sample" / "sample_workbook.xlsx"
    comp = build(out)
    print("wrote %s (%d companies, %d core names, all invented)" % (out, len(comp), int(comp["is_core"].sum())))
    print("try it:  SCREENER_DATA_FILE=%s .venv/bin/streamlit run app.py" % out)


if __name__ == "__main__":
    main()
