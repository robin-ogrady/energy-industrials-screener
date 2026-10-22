"""
loader.py: read the Bloomberg workbook ONCE, clean it, and cache the result.

What it produces (a dict, see load_data()):
    universe, settings, snapshot, data_dictionary, news   -> tidy tables
    prices, fundamentals, estimates, sentiment, commodities -> long tables
                                       (ticker, date, field, value)
    links            -> one row per main company + related company + direction
    link_review      -> fuzzy / uncertain / unmatched related companies
    quality          -> data quality tables (missing data, broken fields, ...)

Run it directly to (re)build the cache and print a check of the data:
    .venv/bin/python loader.py            # uses the cache if it is fresh
    .venv/bin/python loader.py --force    # re-reads the workbook
"""
import pickle
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from rapidfuzz import fuzz, process

import config

LCFG = config.CONFIG["loader"]
MCFG = config.CONFIG["matching"]

def cache_file_for(path):
    """One cache file per workbook, so switching files (e.g. to the sample)
    never mixes data."""
    return config.CACHE_DIR / ("clean_data_%s.pkl" % Path(path).stem)


CACHE_FILE = cache_file_for(config.data_path())   # cache of the configured workbook
# Bump this when the parsing logic changes, so old caches get rebuilt.
CACHE_VERSION = 2

TABLE_SHEETS = ["Universe", "Settings", "Snapshot", "Links", "News", "Data_Dictionary"]
TS_SHEETS = ["Prices_D", "Fundamentals_Q", "Estimates_W", "Sentiment_D", "Commodities_D"]
SC_TABS = ["SC_Suppliers", "SC_Suppliers_Rev", "SC_Customers"]

# Snapshot columns that are text, not numbers.
SNAPSHOT_TEXT_COLS = ["ticker", "asof", "NAME", "GICS_SECTOR_NAME", "GICS_INDUSTRY_NAME"]

# Supply chain headers (read by NAME, because the tabs order them differently)
# mapped to our own column names. The percent columns differ by direction.
SC_COMMON = {
    "ticker": "related_ticker",
    "company name": "related_name",
    "cost category": "cost_type",
    "relationship value (q) (mln) (usd)": "relationship_value_usd_mm",
    "relationship year": "relationship_year",
    "relationship period": "relationship_period",
    "source": "source",
    "as of date": "as_of_date",
}
SC_SUPPLIER_PCTS = {"%cost category": "exposure_pct", "%supplier rev": "reverse_pct"}
SC_CUSTOMER_PCTS = {"%company rev": "exposure_pct", "%customer cost category": "reverse_pct"}


# ===========================================================================
# 1. Cell cleaning
# ===========================================================================
def clean_value(v):
    """Turn one raw cell into a clean Python value.

    Bloomberg errors ("#N/A N/A", "#N/A Invalid Field", ...) and "--" become
    NaN. Text is stripped of spaces. Numbers and dates pass through unchanged.
    We never invent a number: anything unreadable becomes NaN.
    """
    if v is None:
        return np.nan
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("#") or s in LCFG["missing_strings"]:
            return np.nan
        return s
    if isinstance(v, float) and np.isnan(v):
        return np.nan
    return v


def is_text(v):
    """True for a non-empty string (used to spot real tickers and headers)."""
    return isinstance(v, str) and v != ""


def sheet_rows(wb, name):
    """Read every row of a sheet as a list of tuples of CLEAN values.

    Rows are padded to the same length so column positions line up.
    A missing sheet returns an empty list instead of crashing.
    """
    if name not in wb.sheetnames:
        return []
    rows = [tuple(clean_value(v) for v in r) for r in wb[name].iter_rows(values_only=True)]
    width = max((len(r) for r in rows), default=0)
    return [r + (np.nan,) * (width - len(r)) for r in rows]


# ===========================================================================
# 2. Table sheets (header in row 1, data from row 2)
# ===========================================================================
def parse_table(rows):
    """Turn a header-in-row-1 sheet into a DataFrame.

    Only the first run of named header cells is used, so side notes written
    further right in row 1 (the Universe and Links sheets have these) are
    ignored. Fully empty rows are dropped.
    """
    if not rows:
        return pd.DataFrame()
    header = []
    for h in rows[0]:
        if not is_text(h):
            break
        header.append(h)
    if not header:
        return pd.DataFrame()
    body = [r[: len(header)] for r in rows[1:]]
    df = pd.DataFrame(body, columns=header)
    return df.dropna(how="all").reset_index(drop=True)


def is_example_row(row):
    """True if any text cell in the row says EXAMPLE (template placeholder)."""
    return any(isinstance(v, str) and "EXAMPLE" in v.upper() for v in row.values)


def drop_example_rows(df):
    """Remove template rows such as 'EXAMPLE ROW - delete'."""
    if df.empty:
        return df
    mask = df.apply(is_example_row, axis=1)
    return df[~mask].reset_index(drop=True)


def clean_universe(df):
    """Keep rows with a real ticker and add an is_core flag (demo_set = Y)."""
    if df.empty or "ticker" not in df.columns:
        return pd.DataFrame(columns=["ticker", "chain", "role", "demo_set", "is_core", "universe_row"])
    df = df.copy()
    # Excel row number (header is row 1), useful for checking block order.
    df["universe_row"] = df.index + 2
    df = df[df["ticker"].apply(is_text)].copy()
    df["is_core"] = df.get("demo_set", pd.Series(index=df.index, dtype=object)).astype(str).str.upper().eq("Y")
    return df.reset_index(drop=True)


def clean_snapshot(df):
    """Make Snapshot numeric, mark excluded tickers, handle dividend yield.

    Returns (snapshot, raw_snapshot). raw_snapshot is the numeric version
    BEFORE any assumption (used for the missing-data report).
    """
    if df.empty or "ticker" not in df.columns:
        empty = pd.DataFrame(columns=SNAPSHOT_TEXT_COLS)
        return empty, empty
    df = df[df["ticker"].apply(is_text)].copy()
    for c in df.columns:
        if c not in SNAPSHOT_TEXT_COLS:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    if "asof" in df.columns:
        df["asof"] = pd.to_datetime(df["asof"], errors="coerce")
    df = df.drop_duplicates("ticker").reset_index(drop=True)
    raw = df.copy()

    # Tickers with no price are probably acquired, merged or renamed.
    px = df["PX_LAST"] if "PX_LAST" in df.columns else pd.Series(np.nan, index=df.index)
    df["excluded"] = px.isna()
    df["exclude_reason"] = np.where(df["excluded"], "no price (probably acquired, merged or renamed)", "")

    # Missing dividend yield on a priced stock is most likely a non-payer:
    # set to 0 and flag it. Unpriced (excluded) tickers stay missing.
    if "EQY_DVD_YLD_IND" in df.columns:
        assume = df["EQY_DVD_YLD_IND"].isna() & ~df["excluded"]
        df["div_assumed_zero"] = assume
        df.loc[assume, "EQY_DVD_YLD_IND"] = 0.0
    else:
        df["div_assumed_zero"] = False
    # PE_RATIO is deliberately NOT filled (missing usually = negative earnings).
    return df, raw


# ===========================================================================
# 3. Time series sheets (one column group per company, side by side)
# ===========================================================================
def parse_wide(rows, width, sheet=""):
    """Turn a side-by-side sheet into one long DataFrame.

    Layout: group i starts at column i * width (0-based here). Row 1 has the
    ticker in the group's first column, row 2 the headers ("date", fields),
    data from row 3. Each group has its own dates.

    Returns (long_df, groups_df):
      long_df   columns ticker, date, field, value (only real values)
      groups_df one row per group slot with counts, for the quality report
    """
    cols = ["ticker", "date", "field", "value"]
    if len(rows) < 2 or width < 2:
        return pd.DataFrame(columns=cols), pd.DataFrame()
    n_cols = len(rows[0])
    pieces, groups = [], []
    for i, start in enumerate(range(0, n_cols, width)):
        ticker = rows[0][start]
        info = {"sheet": sheet, "group": i, "ticker": ticker if is_text(ticker) else np.nan}
        if not is_text(ticker):
            info["status"] = "blank slot"
            groups.append(info)
            continue
        headers = list(rows[1][start:start + width])
        # Name any unlabeled column by position so nothing breaks.
        headers = [h if is_text(h) else "col%d" % j for j, h in enumerate(headers)]
        headers[0] = "date"
        block = pd.DataFrame([r[start:start + width] for r in rows[2:]], columns=headers)
        block["date"] = pd.to_datetime(block["date"], errors="coerce")
        block = block[block["date"].notna()]
        long = block.melt(id_vars="date", var_name="field", value_name="value")
        long["value"] = pd.to_numeric(long["value"], errors="coerce")
        info["n_dates"] = len(block)
        for f in headers[1:]:
            info["n_" + f] = int(long.loc[long["field"] == f, "value"].notna().sum())
        long = long.dropna(subset=["value"])
        info["status"] = "ok" if len(long) else "no data"
        if len(long):
            info["first_date"] = long["date"].min()
            info["last_date"] = long["date"].max()
            long.insert(0, "ticker", ticker)
            pieces.append(long[cols])
        groups.append(info)
    out = pd.concat(pieces, ignore_index=True) if pieces else pd.DataFrame(columns=cols)
    return out, pd.DataFrame(groups)


# ===========================================================================
# 4. Supply chain tabs (one 23-row block per core name)
# ===========================================================================
def norm_header(h):
    """Lower-case a header and collapse spaces so lookups are forgiving."""
    return re.sub(r"\s+", " ", str(h).strip().lower())


def parse_sc_tab(rows, tab, direction, core_tickers):
    """Read one SC_ tab into a list of link records.

    Block i starts at row i * 23 (0-based here). Columns are found by header
    NAME. Rows with neither a ticker nor a name are empty slots and skipped.
    Percent columns are converted from percent units to fractions (/100).

    Returns (records_df, blocks_df). blocks_df has one row per block with the
    main ticker found, the ticker expected from Universe, and the row count.
    """
    h = LCFG["sc_block_height"]
    pct_map = SC_SUPPLIER_PCTS if direction == "Supplier" else SC_CUSTOMER_PCTS
    colmap = dict(SC_COMMON, **pct_map)
    records, blocks = [], []
    n_blocks = LCFG["sc_n_blocks"] or len(core_tickers)
    for i in range(n_blocks):
        top = i * h
        expected = core_tickers[i] if i < len(core_tickers) else np.nan
        info = {"tab": tab, "block": i, "expected_main": expected}
        if top + 1 >= len(rows):
            info.update(main_ticker=np.nan, n_rows=0, status="missing block")
            blocks.append(info)
            continue
        main = rows[top][0]
        info["main_ticker"] = main if is_text(main) else np.nan
        headers = [norm_header(x) if is_text(x) else None for x in rows[top + 1]]
        pos = {colmap[x]: j for j, x in enumerate(headers) if x in colmap}
        n = 0
        for r in rows[top + 2: top + 2 + LCFG["sc_max_rows"]]:
            rec = {name: r[j] for name, j in pos.items()}
            if not is_text(rec.get("related_ticker")) and not is_text(rec.get("related_name")):
                continue  # empty slot (fewer than 20 relationships)
            rec.update(main_ticker=info["main_ticker"], direction=direction, tab=tab, rank_in_tab=n + 1)
            records.append(rec)
            n += 1
        info["n_rows"] = n
        info["status"] = "ok" if n else "empty block"
        info["missing_headers"] = ", ".join(sorted(set(colmap.values()) - set(pos)))
        blocks.append(info)

    df = pd.DataFrame(records)
    if not df.empty:
        for c in ["exposure_pct", "reverse_pct"]:
            if c in df.columns:
                # Percent units -> fractions: 3.97 means 3.97 percent -> 0.0397
                df[c] = pd.to_numeric(df[c], errors="coerce") / 100.0
        if "relationship_value_usd_mm" in df.columns:
            df["relationship_value_usd_mm"] = pd.to_numeric(df["relationship_value_usd_mm"], errors="coerce")
    return df, pd.DataFrame(blocks)


def related_key(row):
    """Identity of a related company: its ticker if valid, else its name."""
    t = row.get("related_ticker")
    if is_text(t):
        return t.upper()
    return "NAME:" + str(row.get("related_name", "")).upper()


def merge_suppliers(by_value, by_rev):
    """Combine the two supplier tabs into one row per main company + supplier.

    Same company in both tabs -> one row, values taken from SC_Suppliers
    first and filled from SC_Suppliers_Rev. 'found_in' records which tab(s)
    the link came from. Returns (merged, n_conflicts).
    """
    frames = [d for d in (by_value, by_rev) if not d.empty]
    if not frames:
        return pd.DataFrame(), 0
    both = pd.concat(frames, ignore_index=True)
    both["related_key"] = both.apply(related_key, axis=1)
    value_cols = [c for c in both.columns if c not in ("main_ticker", "related_key", "tab", "rank_in_tab", "direction")]
    rows, conflicts = [], 0
    for (main, key), g in both.groupby(["main_ticker", "related_key"], sort=False):
        g = g.sort_values("tab")  # "SC_Suppliers" sorts before "SC_Suppliers_Rev"
        rec = {"main_ticker": main, "related_key": key, "direction": "Supplier"}
        for c in value_cols:
            vals = g[c].dropna()
            rec[c] = vals.iloc[0] if len(vals) else np.nan
            if c in ("exposure_pct", "reverse_pct", "relationship_value_usd_mm") and len(vals) > 1:
                if (vals.max() - vals.min()) > 1e-9:
                    conflicts += 1
        tabs = set(g["tab"])
        rec["found_in"] = "both" if len(tabs) > 1 else ("by value" if "SC_Suppliers" in tabs else "by supplier rev %")
        rows.append(rec)
    return pd.DataFrame(rows), conflicts


# ---- name matching ----------------------------------------------------------
# Words that do not help identify a company. Removing them lets
# "Exxon Mobil Corp" match "EXXON MOBIL CORP".
NAME_STOPWORDS = {
    "INC", "INCORPORATED", "CORP", "CORPORATION", "CO", "COMPANY", "LTD", "LIMITED",
    "PLC", "SA", "AG", "NV", "LLC", "LP", "THE", "HOLDINGS", "HOLDING", "HLDGS",
    "GROUP", "GRP", "ADR", "CLASS", "CL", "A", "B", "DE", "SE", "ASA", "AB", "BHD",
    "TBK", "PT", "KK", "SPA", "CV", "SAB", "OYJ",
}


def norm_name(name):
    """Upper-case, strip punctuation and legal-form words from a name."""
    s = re.sub(r"[^A-Z0-9 ]", " ", str(name).upper().replace("&", " AND "))
    return " ".join(w for w in s.split() if w not in NAME_STOPWORDS)


def match_related(links, snapshot):
    """Decide, for each related company, whether it is in our universe.

    Step 1: exact ticker match with a Snapshot ticker.
    Step 2: fuzzy name match against Snapshot NAME (rapidfuzz). Accepted only
            at or above fuzzy_accept. A related company that already has a
            different US ticker is never auto-accepted (two US tickers are
            normally two different securities), it goes to review instead.
    We never force a match. Returns (links_with_match_cols, review_df).
    """
    links = links.copy()
    snap_tickers = set(snapshot["ticker"]) if not snapshot.empty else set()
    names = {}
    if not snapshot.empty and "NAME" in snapshot.columns:
        names = {t: norm_name(n) for t, n in zip(snapshot["ticker"], snapshot["NAME"]) if is_text(n)}
    choices = list(names.values())
    choice_tickers = list(names.keys())

    uniq = links.drop_duplicates("related_key")[["related_key", "related_ticker", "related_name"]]
    result, review = {}, []
    for _, r in uniq.iterrows():
        t, n = r["related_ticker"], r["related_name"]
        if is_text(t) and t in snap_tickers:
            result[r["related_key"]] = (t, "exact ticker", 100.0)
            continue
        best_t, score = np.nan, 0.0
        if is_text(n) and choices:
            hit = process.extractOne(norm_name(n), choices, scorer=fuzz.token_sort_ratio)
            if hit:
                best_t, score = choice_tickers[hit[2]], float(hit[1])
        us_ticker = is_text(t) and t.upper().endswith(" US EQUITY")
        if score >= MCFG["fuzzy_accept"] and not us_ticker:
            status, matched = "fuzzy name (accepted)", best_t
        elif score >= MCFG["fuzzy_review"]:
            status, matched = "uncertain (not matched)", np.nan
        else:
            status, matched = "unmatched", np.nan
        result[r["related_key"]] = (matched, status, score)
        review.append({
            "related_ticker": t, "related_name": n, "status": status,
            "best_candidate": best_t, "candidate_name": snapshot.set_index("ticker")["NAME"].get(best_t) if is_text(best_t) else np.nan,
            "name_score": round(score, 1),
        })
    links["universe_ticker"] = links["related_key"].map(lambda k: result[k][0])
    links["match_method"] = links["related_key"].map(lambda k: result[k][1])
    links["in_universe"] = links["universe_ticker"].apply(is_text)
    review = pd.DataFrame(review)
    if not review.empty:
        order = {"fuzzy name (accepted)": 0, "uncertain (not matched)": 1, "unmatched": 2}
        review = review.sort_values(["status", "name_score"], key=lambda s: s.map(order) if s.name == "status" else -s)
    return links, review


def build_links(sc_frames, snapshot):
    """Combine the three SC tabs into the final links table.

    Columns follow the spec: main_ticker, direction, related_name,
    related_ticker, exposure_pct, reverse_pct, cost_type,
    relationship_value_usd_mm, source, as_of_date (plus helper columns).
    """
    sup, conflicts = merge_suppliers(sc_frames.get("SC_Suppliers", pd.DataFrame()),
                                     sc_frames.get("SC_Suppliers_Rev", pd.DataFrame()))
    cust = sc_frames.get("SC_Customers", pd.DataFrame()).copy()
    if not cust.empty:
        cust["related_key"] = cust.apply(related_key, axis=1)
        cust["found_in"] = "customers"
        cust = cust.drop(columns=["tab", "rank_in_tab"], errors="ignore")
        # The same customer twice in one block would double count it.
        cust = cust.drop_duplicates(["main_ticker", "related_key"])
    links = pd.concat([d for d in (sup, cust) if not d.empty], ignore_index=True) if (not sup.empty or not cust.empty) else pd.DataFrame()
    if links.empty:
        return links, pd.DataFrame(), conflicts

    links["cost_type"] = links["cost_type"].where(links["cost_type"].apply(is_text))
    links["cost_type"] = links["cost_type"].str.upper()
    links["cost_type_missing"] = links["cost_type"].isna()
    links["source"] = links["source"].astype(object)
    links["source_type"] = np.where(links["source"].astype(str).str.lower().str.startswith("estimate"), "Estimate",
                                    np.where(links["source"].apply(is_text), "Reported", "Unknown"))
    links["as_of_date"] = pd.to_datetime(links["as_of_date"], errors="coerce", format="%m/%d/%Y")
    links["as_of_year"] = links["as_of_date"].dt.year
    links["relationship_year"] = pd.to_numeric(links["relationship_year"], errors="coerce")

    links, review = match_related(links, snapshot)
    front = ["main_ticker", "direction", "related_name", "related_ticker", "exposure_pct", "reverse_pct",
             "cost_type", "relationship_value_usd_mm", "source", "as_of_date"]
    rest = [c for c in links.columns if c not in front]
    return links[front + rest], review, conflicts


# ===========================================================================
# 5. Data quality report
# ===========================================================================
def quality_tables(d, raw_snapshot, ts_groups, sc_blocks):
    """Build the data quality tables shown on the Data quality page."""
    q = {}
    uni = d["universe"]
    snap_fields = [c for c in raw_snapshot.columns if c not in ("ticker", "asof")]

    # Missing percent per Snapshot field (over all universe tickers, raw values).
    miss = raw_snapshot[snap_fields].isna().mean().rename("pct_missing").to_frame()
    miss["sheet"] = "Snapshot"
    q["snapshot_field_missing"] = miss.reset_index().rename(columns={"index": "field"})

    # Missing percent per ticker (Snapshot), excluding the fully broken fields.
    ok_fields = [f for f in snap_fields if f not in d["broken_fields"]]
    per_t = raw_snapshot.set_index("ticker")[ok_fields].isna().mean(axis=1).rename("pct_missing_snapshot")
    q["ticker_missing"] = per_t.reset_index()

    # Time series: per sheet and field, share of companies with no values at all,
    # and the number of values per company.
    rows = []
    for sheet, g in ts_groups.items():
        if g.empty:
            continue
        real = g[g["status"] != "blank slot"]
        for c in [c for c in g.columns if c.startswith("n_") and c != "n_dates"]:
            vals = real[c].fillna(0)
            rows.append({
                "sheet": sheet, "field": c[2:], "companies": len(real),
                "pct_companies_no_data": float((vals == 0).mean()) if len(real) else np.nan,
                "median_values_per_company": float(vals.median()) if len(real) else np.nan,
            })
    q["ts_field_coverage"] = pd.DataFrame(rows)
    q["ts_groups"] = pd.concat([g for g in ts_groups.values() if not g.empty], ignore_index=True) \
        if any(not g.empty for g in ts_groups.values()) else pd.DataFrame()

    # Time series values per ticker per sheet (wide table, for the heatmap).
    if not q["ts_groups"].empty:
        tg = q["ts_groups"][q["ts_groups"]["status"] != "blank slot"]
        q["ticker_ts_counts"] = tg.pivot_table(index="ticker", columns="sheet", values="n_dates", aggfunc="first")

    # Data_Dictionary + measured missing percent + "working" flag.
    dd = d["data_dictionary"].copy()
    if not dd.empty and "column_name" in dd.columns:
        miss_map = dict(zip(q["snapshot_field_missing"]["field"], q["snapshot_field_missing"]["pct_missing"]))
        dd["pct_missing_snapshot"] = dd["column_name"].map(miss_map)
        dd["working"] = dd["pct_missing_snapshot"] < LCFG["working_field_max_missing"]
        dd["broken"] = dd["column_name"].isin(d["broken_fields"])
    q["dictionary"] = dd

    snap = d["snapshot"]
    q["excluded"] = snap.loc[snap["excluded"], ["ticker", "NAME", "exclude_reason"]].reset_index(drop=True)
    q["div_assumed_zero"] = snap.loc[snap["div_assumed_zero"], ["ticker", "NAME"]].reset_index(drop=True)
    q["sector_counts"] = snap["GICS_SECTOR_NAME"].fillna("(missing)").value_counts().rename_axis("sector").reset_index(name="companies")
    q["sc_blocks"] = sc_blocks

    links = d["links"]
    core = uni.loc[uni["is_core"], "ticker"].tolist()
    if not links.empty:
        n_sup = links[links["direction"] == "Supplier"].groupby("main_ticker").size()
        n_cus = links[links["direction"] == "Customer"].groupby("main_ticker").size()
    else:
        n_sup = n_cus = pd.Series(dtype=int)
    q["core_link_counts"] = pd.DataFrame({
        "ticker": core,
        "suppliers": [int(n_sup.get(t, 0)) for t in core],
        "customers": [int(n_cus.get(t, 0)) for t in core],
    })
    return q


# ===========================================================================
# 6. Main entry point with caching
# ===========================================================================
def source_signature(path):
    """File name + size + modified time. A change means rebuild the cache."""
    st = path.stat()
    return (path.name, st.st_size, int(st.st_mtime))


def parse_workbook(path):
    """Read the workbook and build every clean table. Slow part (~seconds)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    d = {}
    raw = {name: parse_table(sheet_rows(wb, name)) for name in TABLE_SHEETS}

    d["universe"] = clean_universe(raw["Universe"])
    d["settings"] = raw["Settings"]
    snap, raw_snap = clean_snapshot(raw["Snapshot"])

    # Drop Snapshot fields that are 100 percent missing among real tickers.
    numeric = [c for c in raw_snap.columns if c not in SNAPSHOT_TEXT_COLS]
    d["broken_fields"] = [c for c in numeric if raw_snap[c].isna().all()] if len(raw_snap) else []
    d["snapshot"] = snap.drop(columns=d["broken_fields"])
    d["data_dictionary"] = raw["Data_Dictionary"]

    news = drop_example_rows(raw["News"])
    if not news.empty:
        news["date"] = pd.to_datetime(news.get("date"), errors="coerce")
        news["manual_score"] = pd.to_numeric(news.get("manual_score"), errors="coerce")
    d["news"] = news
    d["links_sheet"] = drop_example_rows(raw["Links"])  # spec: ignored, kept for the report

    ts_groups = {}
    names = {"Prices_D": "prices", "Fundamentals_Q": "fundamentals", "Estimates_W": "estimates",
             "Sentiment_D": "sentiment", "Commodities_D": "commodities"}
    for sheet in TS_SHEETS:
        long, groups = parse_wide(sheet_rows(wb, sheet), LCFG["group_width"][sheet], sheet)
        d[names[sheet]] = long
        ts_groups[sheet] = groups

    core = d["universe"].loc[d["universe"]["is_core"], "ticker"].tolist()
    sc_frames, sc_blocks = {}, []
    for tab in SC_TABS:
        direction = "Customer" if tab == "SC_Customers" else "Supplier"
        recs, blocks = parse_sc_tab(sheet_rows(wb, tab), tab, direction, core)
        sc_frames[tab] = recs
        sc_blocks.append(blocks)
    d["sc_raw"] = sc_frames
    d["links"], d["link_review"], d["supplier_conflicts"] = build_links(sc_frames, d["snapshot"])
    d["quality"] = quality_tables(d, raw_snap, ts_groups, pd.concat(sc_blocks, ignore_index=True))
    d["raw_snapshot"] = raw_snap
    wb.close()
    return d


def load_data(force=False, path=None):
    """Return the clean data dict, from the cache when it is still fresh."""
    path = Path(path or config.data_path())
    if not path.exists():
        raise FileNotFoundError(
            "Data file not found: %s\nSet DATA_FILE in config.py or the SCREENER_DATA_FILE environment variable. "
            "To try the tool without Bloomberg, run: .venv/bin/python make_sample_data.py" % path)
    sig = source_signature(path)
    cache_file = cache_file_for(path)
    if cache_file.exists() and not force:
        with open(cache_file, "rb") as f:
            cached = pickle.load(f)
        if cached.get("_signature") == sig and cached.get("_version") == CACHE_VERSION:
            return cached
    d = parse_workbook(path)
    d["_signature"], d["_version"] = sig, CACHE_VERSION
    config.CACHE_DIR.mkdir(exist_ok=True)
    with open(cache_file, "wb") as f:
        pickle.dump(d, f)
    # Human-readable copies (all gitignored).
    if not d["links"].empty:
        d["links"].to_csv(config.CACHE_DIR / "links.csv", index=False)
    if not d["link_review"].empty:
        d["link_review"].to_csv(config.CACHE_DIR / "link_match_review.csv", index=False)
    return d


# ===========================================================================
# 7. Data summary printout (python loader.py)
# ===========================================================================
def data_summary(d):
    """What the loader found, plus layout checks that should say OK.

    Run this after every new Bloomberg pull and read it top to bottom:
    rows marked CHECK point at a layout problem in the workbook.
    """
    uni, snap, links, q = d["universe"], d["snapshot"], d["links"], d["quality"]
    out = []

    def info(item, value):
        out.append((item, value, ""))

    def check(item, ok, value=""):
        out.append((item, value, "OK" if ok else "CHECK"))

    short = lambda ts: ", ".join(sorted(str(t).split()[0] for t in ts)) or "none"
    info("Universe tickers", len(uni))
    n_core = int(uni["is_core"].sum())
    check("Core names (demo_set = Y)", n_core > 0, n_core)
    check("Snapshot tickers match Universe", set(snap["ticker"]) == set(uni["ticker"]), len(snap))
    info("Fully broken Snapshot fields (dropped)", ", ".join(d["broken_fields"]) or "none")
    info("Tickers with no price (excluded)", short(snap.loc[snap["excluded"], "ticker"]))
    priced = snap[~snap["excluded"]]
    info("Priced names with no dividend yield (set to 0)", "%d (%.0f%%)" % (
        priced["div_assumed_zero"].sum(), 100 * priced["div_assumed_zero"].mean() if len(priced) else 0))
    info("Sector mix", ", ".join("%s %d" % (r["sector"], r["companies"]) for _, r in q["sector_counts"].iterrows()))

    g = q["ts_groups"]
    order = dict(zip(uni["universe_row"] - 2, uni["ticker"]))
    for sheet in TS_SHEETS:
        gs = g[g["sheet"] == sheet] if len(g) else g
        if gs is None or len(gs) == 0:
            check("%s present" % sheet, False, "missing or empty")
            continue
        real = gs[gs["status"] != "blank slot"]
        info("%s group slots / with ticker / with data" % sheet,
             "%d / %d / %d" % (len(gs), len(real), (gs["status"] == "ok").sum()))
        if sheet != "Commodities_D":
            same = all(order.get(r["group"]) == r["ticker"] for _, r in real.iterrows())
            check("%s group order matches Universe" % sheet, same)
    if len(d["prices"]):
        pr = d["prices"].groupby("ticker").size()
        info("Prices_D rows per company (median / min)", "%d / %d" % (pr.median(), pr.min()))
        info("Prices_D date range", "%s to %s" % (d["prices"]["date"].min().date(), d["prices"]["date"].max().date()))
    if len(d["fundamentals"]):
        info("Fundamentals_Q quarters (median)", int(d["fundamentals"].groupby("ticker")["date"].nunique().median()))
    if len(d["estimates"]):
        info("Estimates_W weeks (median)", int(d["estimates"].groupby("ticker")["date"].nunique().median()))
    if len(d["sentiment"]):
        sd = d["sentiment"].groupby("ticker").size()
        info("Sentiment_D news days per company (max / median)", "%d / %d" % (sd.max(), sd.median()))
    want = set(config.CONFIG["features"]["commodities"])
    have = set(d["commodities"]["ticker"]) if len(d["commodities"]) else set()
    check("Commodity series found", want <= have, ", ".join(sorted(have)) or "none")
    info("News rows (example rows dropped)", len(d["news"]))

    if not links.empty:
        raw = d["sc_raw"]
        sup = links[links["direction"] == "Supplier"]
        cus = links[links["direction"] == "Customer"]
        info("SC_Suppliers / SC_Suppliers_Rev rows", "%d / %d" % (len(raw["SC_Suppliers"]), len(raw["SC_Suppliers_Rev"])))
        info("Supplier links after merging both tabs", len(sup))
        info("Customer links", len(cus))
        check("Core names with supplier links", sup["main_ticker"].nunique() == n_core,
              "%d of %d" % (sup["main_ticker"].nunique(), n_core))
        cl = q["core_link_counts"]
        info("Core names with no customer links", short(cl.loc[cl["customers"] == 0, "ticker"]))
        info("Links with Source = Estimate", "%.0f%%" % (100 * (links["source_type"] == "Estimate").mean()))
        info("Links missing cost category", int(links["cost_type_missing"].sum()))
        uniq = links.drop_duplicates("related_key")
        info("Unique related companies / in universe", "%d / %d" % (len(uniq), uniq["in_universe"].sum()))
        check("Supplier values agree between the two supplier tabs", d["supplier_conflicts"] == 0, d["supplier_conflicts"])
        bl = q["sc_blocks"]
        check("SC blocks follow Universe core order", int((bl["main_ticker"] != bl["expected_main"]).sum()) == 0)
    else:
        info("Supply chain links", "none found")
    return pd.DataFrame(out, columns=["item", "value", "check"])


def write_quality_report(d):
    """Write cache/data_quality_report.txt: a plain-text data quality summary."""
    q = d["quality"]
    pct = lambda x: "%.1f%%" % (100 * x)
    lines = ["DATA QUALITY REPORT", "Data file: %s" % config.DATA_FILE, ""]

    lines.append("1. Snapshot fields, percent missing (all %d universe tickers, before any assumption)" % len(d["universe"]))
    m = q["snapshot_field_missing"].sort_values("pct_missing", ascending=False)
    for _, r in m.iterrows():
        tag = "  BROKEN, dropped" if r["field"] in d["broken_fields"] else ""
        lines.append("   %-26s %7s%s" % (r["field"], pct(r["pct_missing"]), tag))

    lines += ["", "2. Time series sheets"]
    for _, r in q["ts_field_coverage"].iterrows():
        lines.append("   %-15s %-26s companies with no data: %6s   median values per company: %d" % (
            r["sheet"], r["field"], pct(r["pct_companies_no_data"]), r["median_values_per_company"]))

    lines += ["", "3. Excluded tickers (not scored)"]
    for _, r in q["excluded"].iterrows():
        lines.append("   %-16s %-32s %s" % (r["ticker"], r["NAME"], r["exclude_reason"]))

    lines += ["", "4. Tickers missing the most Snapshot fields (broken fields not counted)"]
    tm = q["ticker_missing"].sort_values("pct_missing_snapshot", ascending=False)
    for _, r in tm[tm["pct_missing_snapshot"] > 0.10].iterrows():
        lines.append("   %-16s %s" % (r["ticker"], pct(r["pct_missing_snapshot"])))

    lines += ["", "5. Assumption: missing dividend yield treated as 0 for %d priced companies" % len(q["div_assumed_zero"])]

    lines += ["", "6. Data_Dictionary field_status vs measured missing data (working = under %s missing)"
              % pct(LCFG["working_field_max_missing"])]
    dd = q["dictionary"]
    for _, r in dd.iterrows():
        mp = r.get("pct_missing_snapshot")
        lines.append("   %-40s %-28s missing %-7s %s" % (
            r["column_name"], r["field_status"], pct(mp) if pd.notna(mp) else "n/a",
            "BROKEN" if r["broken"] else ("working" if r["working"] else "")))

    lines += ["", "7. Supply chain"]
    L = d["links"]
    if L.empty:
        lines.append("   No supply chain links found.")
        path = config.CACHE_DIR / "data_quality_report.txt"
        path.write_text("\n".join(lines) + "\n")
        return path
    uniq = L.drop_duplicates("related_key")
    lines.append("   Supplier links (both supplier tabs merged): %d" % (L["direction"] == "Supplier").sum())
    lines.append("   Customer links: %d" % (L["direction"] == "Customer").sum())
    lines.append("   Unique related companies: %d, in universe: %d (%s)" % (
        len(uniq), uniq["in_universe"].sum(), pct(uniq["in_universe"].mean())))
    cl = q["core_link_counts"]
    lines.append("   Core names with no customer links: %s" % ", ".join(t.split()[0] for t in cl.loc[cl["customers"] == 0, "ticker"]))
    lines.append("   Core names with no supplier links: %s" % (", ".join(t.split()[0] for t in cl.loc[cl["suppliers"] == 0, "ticker"]) or "none"))
    lines.append("   Match review list: cache/link_match_review.csv")
    path = config.CACHE_DIR / "data_quality_report.txt"
    path.write_text("\n".join(lines) + "\n")
    return path


if __name__ == "__main__":
    data = load_data(force="--force" in sys.argv)
    pd.set_option("display.width", 200)
    pd.set_option("display.max_colwidth", 90)
    pd.set_option("display.max_rows", 200)
    print("Data file:", config.data_path())
    print(data_summary(data).to_string(index=False))
    print("\nWrote", write_quality_report(data))
